import json
from pathlib import Path

import pytest
from docx import Document

from cv_ats.cli import main
from cv_ats.extract import UnsupportedFormatError, extract_text

ROOT = Path(__file__).resolve().parents[1]
CV_SAMPLE = ROOT / "examples" / "cv_sample.txt"
JD_SAMPLE = ROOT / "examples" / "job_description.txt"


def _write_minimal_pdf(path: Path, lines: list[str]) -> None:
    """Write a tiny but valid one-page PDF with ASCII text, so the PDF path is tested without binaries."""
    text_ops = " ".join(f"({line}) Tj 0 -16 Td" for line in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {text_ops} ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    path.write_bytes(out.encode("ascii"))


def test_reads_txt(tmp_path):
    file = tmp_path / "cv.txt"
    file.write_text("Golang and Postgres", encoding="utf-8")
    assert extract_text(file) == "Golang and Postgres"


def test_reads_docx_paragraphs_and_tables(tmp_path):
    file = tmp_path / "cv.docx"
    document = Document()
    document.add_paragraph("Backend developer")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker Compose"
    document.save(str(file))

    text = extract_text(file)
    assert "Backend developer" in text
    assert "Docker Compose" in text


def test_reads_pdf(tmp_path):
    file = tmp_path / "cv.pdf"
    _write_minimal_pdf(file, ["Python developer", "Kubernetes and Redis"])
    text = extract_text(file)
    assert "Python developer" in text
    assert "Kubernetes" in text


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_text(tmp_path / "nope.txt")


def test_unsupported_extension_raises(tmp_path):
    file = tmp_path / "cv.odt"
    file.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFormatError):
        extract_text(file)


def test_cli_prints_report_for_sample_files(capsys):
    assert main(["check", "--cv", str(CV_SAMPLE), "--jd", str(JD_SAMPLE)]) == 0
    out = capsys.readouterr().out
    assert "Coverage score:" in out
    assert "Go" in out


def test_cli_json_output_is_valid(capsys):
    assert main(["check", "--cv", str(CV_SAMPLE), "--jd", str(JD_SAMPLE), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert "Kubernetes" in data["skills_missing"]
    assert "Go" in data["skills_matched"]


def test_cli_min_score_sets_exit_status(capsys):
    assert main(["check", "--cv", str(CV_SAMPLE), "--jd", str(JD_SAMPLE), "--min-score", "101"]) == 1


def test_cli_missing_file_returns_2(capsys):
    assert main(["check", "--cv", "does-not-exist.txt", "--jd", str(JD_SAMPLE)]) == 2
    assert "No such file" in capsys.readouterr().err
