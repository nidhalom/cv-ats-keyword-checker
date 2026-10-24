"""Read plain text from .txt, .md, .docx and .pdf files."""

from __future__ import annotations

from pathlib import Path

SUPPORTED_SUFFIXES = (".txt", ".md", ".docx", ".pdf")


class UnsupportedFormatError(ValueError):
    """Raised when the file extension is not one of SUPPORTED_SUFFIXES."""


def extract_text(path: str | Path) -> str:
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"No such file: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix in (".txt", ".md"):
        return file_path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix == ".docx":
        return _read_docx(file_path)
    if suffix == ".pdf":
        return _read_pdf(file_path)
    raise UnsupportedFormatError(
        f"Unsupported file type '{suffix or '(none)'}'. Use one of: {', '.join(SUPPORTED_SUFFIXES)}"
    )


def _read_docx(path: Path) -> str:
    from docx import Document  # imported lazily so .txt users do not need python-docx

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # CVs often put skills in tables, so read the cells too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(part for part in parts if part.strip())


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)
